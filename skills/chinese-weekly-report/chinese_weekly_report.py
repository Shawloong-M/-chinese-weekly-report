#!/usr/bin/env python3
"""Chinese Weekly/Monthly Report Generator v2.3"""
import json, sys, os
from pathlib import Path
from datetime import datetime, timedelta

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import CategoryChartData
    from pptx.oxml.ns import qn
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    from pptx.dml.color import RGBColor
except ImportError:
    RGBColor = None

C = {
    "primary": RGBColor(0x1A, 0x3C, 0x6E),
    "secondary": RGBColor(0x2B, 0x5C, 0x9A),
    "accent": RGBColor(0x3A, 0x7B, 0xD5),
    "green": RGBColor(0x27, 0xAE, 0x60),
    "red": RGBColor(0xE7, 0x4C, 0x3C),
    "orange": RGBColor(0xE8, 0x6C, 0x00),
    "gold": RGBColor(0xF3, 0x9C, 0x12),
    "light_bg": RGBColor(0xF0, 0xF4, 0xF8),
    "dark": RGBColor(0x33, 0x33, 0x33),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
    "border": RGBColor(0xCC, 0xCC, 0xCC),
    "cover_bg1": RGBColor(0x0F, 0x2B, 0x55),
    "cover_bg2": RGBColor(0x1A, 0x3C, 0x6E),
    "cover_accent": RGBColor(0x3A, 0x7B, 0xD5),
}
W, H = 13.333, 7.5
ML, MR = 0.35, 0.35
CW = W - ML - MR
SH = 0.5
ST = 0.6
SHH = 0.45

FONT = {
    "title": 14, "summary": 10, "section": 12, "sub": 10,
    "th": 9, "td": 8, "kn": 9, "kv": 18, "kt": 7,
    "desc": 8, "path": 7, "gt": 10, "gn": 8, "gd": 7, "logo": 9,
}

# ── Helpers ──

def _ea(run, name):
    rPr = run._r.get_or_add_rPr()
    rPr.set(qn('a:ea'), name)

def tb(slide, l, t, w, h, text, sz=None, bold=False, color=None,
       align=PP_ALIGN.LEFT, fn="等线", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    if anchor:
        try: tf.vertical_anchor = anchor
        except: pass
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(sz or FONT["desc"])
    p.font.bold = bold
    p.font.color.rgb = color or C["dark"]
    p.font.name = fn
    for r in p.runs:
        _ea(r, fn)
    p.alignment = align
    return box

def rtb(slide, l, t, w, h, runs, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    try: tf.vertical_anchor = MSO_ANCHOR.TOP
    except: pass
    p = tf.paragraphs[0]
    for i, (txt, sz, bld, clr, fn) in enumerate(runs):
        run = p.add_run() if i > 0 else (p.runs[0] if p.runs else p.add_run())
        run.text = txt
        run.font.size = Pt(sz)
        run.font.bold = bld
        run.font.color.rgb = clr or C["dark"]
        run.font.name = fn
        _ea(run, fn)
    p.alignment = align
    return box

def rect(slide, l, t, w, h, fill, text="", sz=None, fc=None, bold=False,
         fn="等线", bc=None, bw=0):
    shape = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if bc:
        shape.line.color.rgb = bc
        shape.line.width = Pt(bw)
    else:
        shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.08)
        tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.03)
        tf.margin_bottom = Inches(0.03)
        try: tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        except: pass
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(sz or FONT["desc"])
        p.font.bold = bold
        p.font.color.rgb = fc or C["white"]
        p.font.name = fn
        p.alignment = PP_ALIGN.CENTER
    return shape

def logo(slide, text="LOGO"):
    rect(slide, 0.2, 0.07, 0.75, 0.36, C["accent"], text, FONT["logo"], C["white"], True)

def header(slide, title, logo_text="LOGO"):
    rect(slide, 0, 0, W, SH, C["primary"], "")
    logo(slide, logo_text)
    tb(slide, 1.15, 0.06, 10, 0.38, title, FONT["title"], True, C["white"])

def summary_bar(slide, top, text):
    rect(slide, ML, top, CW, SHH, C["light_bg"], "", bc=C["border"], bw=0.5)
    tb(slide, ML + 0.1, top + 0.03, CW - 0.2, SHH - 0.06, text, FONT["summary"], color=C["dark"])

# ── Table (B: improved header color, larger fonts, light_bg alt rows) ──
def add_table(slide, l, t, w, h, headers, rows, cw=None):
    nr = len(rows) + 1
    nc = len(headers)
    ts = slide.shapes.add_table(nr, nc, Inches(l), Inches(t), Inches(w), Inches(h))
    tbl = ts.table
    if cw:
        for i, c in enumerate(cw):
            if i < len(tbl.columns):
                tbl.columns[i].width = Inches(c)
    for ci, hd in enumerate(headers):
        c = tbl.cell(0, ci)
        c.text = hd
        c.fill.solid()
        c.fill.fore_color.rgb = C["secondary"]
        for p in c.text_frame.paragraphs:
            p.font.size = Pt(FONT["th"])
            p.font.bold = True
            p.font.color.rgb = C["white"]
            p.font.name = "等线"
            p.alignment = PP_ALIGN.CENTER
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri + 1, ci)
            c.text = str(val) if val is not None else ""
            c.fill.solid()
            c.fill.fore_color.rgb = C["white"] if ri % 2 == 0 else C["light_bg"]
            for p in c.text_frame.paragraphs:
                p.font.size = Pt(FONT["td"])
                p.font.color.rgb = C["dark"]
                p.font.name = "等线"
                p.alignment = PP_ALIGN.CENTER

# ── Gantt (D: larger name area, thicker bars, grid background) ──
def gantt(slide, l, t, w, h, gd):
    if not gd or not gd.get("tasks"):
        return
    tasks = gd["tasks"]
    ss, es = gd.get("start_date", ""), gd.get("end_date", "")
    try:
        sd = datetime.strptime(ss, "%Y-%m-%d") if ss else datetime.now()
        ed = datetime.strptime(es, "%Y-%m-%d") if es else (sd + timedelta(days=30))
    except:
        sd, ed = datetime.now(), datetime.now() + timedelta(days=30)
    nd = max(1, (ed - sd).days)
    tb(slide, l, t, w, 0.28, "项目进度甘特图", FONT["gt"], True, C["primary"])
    ct = t + 0.32
    ch = h - 0.38
    nw = 1.5
    ba = l + nw
    bw = w - nw
    nt = min(nd, 14)
    step = nd / nt
    for i in range(nt):
        td = sd + timedelta(days=int(i * step))
        x = ba + (i * step / nd) * bw
        tw = (step / nd) * bw
        tb(slide, x, ct - 0.05, tw, 0.22, td.strftime("%m/%d"), FONT["gd"], color=C["dark"], align=PP_ALIGN.CENTER)
    row_h = min(ch / max(len(tasks), 1), 0.42)
    bar_h = max(row_h * 0.6, 0.16)
    for ti, task in enumerate(tasks):
        y = ct + 0.2 + ti * row_h
        nm = task.get("name", f"任务{ti+1}")
        ts_, te_ = task.get("start", ""), task.get("end", "")
        chx = task.get("color", "#3A7BD5")
        try:
            clr = RGBColor(int(chx[1:3], 16), int(chx[3:5], 16), int(chx[5:7], 16))
        except:
            clr = C["accent"]
        # Subtle row background
        rect(slide, ba, y, bw, row_h - 0.02, C["light_bg"] if ti % 2 == 0 else C["white"], "", bc=RGBColor(0xE0,0xE0,0xE0), bw=0.3)
        tb(slide, l, y + 0.02, nw - 0.05, bar_h + 0.05, nm, FONT["gn"], color=C["dark"])
        try:
            t1 = datetime.strptime(ts_, "%Y-%m-%d") if ts_ else sd
            t2 = datetime.strptime(te_, "%Y-%m-%d") if te_ else ed
        except:
            t1, t2 = sd, ed
        bl = ba + max(0, (t1 - sd).days / nd) * bw
        bw2 = max(0.15, max(1, (t2 - t1).days) / nd * bw)
        rect(slide, bl, y + 0.03, bw2, bar_h, clr, "", bc=RGBColor(0x99,0x99,0x99), bw=0.5)

def add_chart(slide, l, t, w, h, spec):
    m = {"line": XL_CHART_TYPE.LINE_MARKERS, "bar": XL_CHART_TYPE.COLUMN_CLUSTERED}
    cd = CategoryChartData()
    cd.categories = spec.get("categories", [])
    for s in spec.get("series", []):
        cd.add_series(s.get("name", ""), s.get("values", []))
    cf = slide.shapes.add_chart(m.get(spec.get("type","bar"), XL_CHART_TYPE.COLUMN_CLUSTERED), Inches(l), Inches(t), Inches(w), Inches(h), cd)
    ch = cf.chart
    ch.has_legend = True
    ch.legend.position = 2
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(8)
    pl = ch.plots[0]
    pl.has_data_labels = True
    pl.data_labels.font.size = Pt(8)
    pl.data_labels.show_value = True

# ══════════════════════════════════════════════════════════════════════════
# Slides
# ══════════════════════════════════════════════════════════════════════════

def s_cover(prs, d):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.shapes.add_shape(1, Inches(0), Inches(0), Inches(W), Inches(H))
    bg.fill.solid(); bg.fill.fore_color.rgb = C["cover_bg1"]; bg.line.fill.background()
    rect(sl, 0, H-1.2, W, 1.2, C["cover_bg2"])
    rect(sl, 2.5, 3.0, 8.333, 0.05, C["cover_accent"])
    for x, w, h in [(0.5,2.5,1.5),(3.5,2.0,1.8),(6.0,3.0,1.2),(9.5,1.8,1.6),(11.0,2.0,1.0)]:
        rect(sl, x, H-1.2-h, w, h, C["cover_bg2"])
        for ox in [0.3, 0.8]: rect(sl, x+ox, H-1.2-h+0.2, 0.3, 0.2, C["gold"])
        if w > 2: rect(sl, x+1.3, H-1.2-h+0.2, 0.3, 0.2, C["gold"])
    rect(sl, 0.5, 0.35, 1.1, 0.45, C["accent"], d.get("logo_text","LOGO"), 12, C["white"], True)
    rt = d.get("report_type","weekly"); pd = d.get("report_period","报告周期")
    tl = {"weekly":"周报","monthly":"月报","quarterly":"季报"}
    tb(sl, 1.5, 1.3, 10.3, 1.0, f"{pd}  工作{tl.get(rt,'报告')}", 36, True, C["white"], align=PP_ALIGN.CENTER)
    pts = [f"部门：{d.get('department','')}"] if d.get("department") else []
    if d.get("author"): pts.append(f"报告人：{d['author']}")
    if d.get("date"): pts.append(f"日期：{d['date']}")
    if pts: tb(sl, 1.5, 3.3, 10.3, 0.5, "  |  ".join(pts), 14, color=RGBColor(0xBB,0xCC,0xDD), align=PP_ALIGN.CENTER)
    tb(sl, 1.5, H-0.5, 10.3, 0.35, "CONFIDENTIAL  ·  数字化推进部  ·  MES项目组", 9, color=RGBColor(0x88,0x99,0xAA), align=PP_ALIGN.CENTER)

def s_toc(prs, d):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "目  录", d.get("logo_text","LOGO"))
    items = [("01","本周概述","工作总结与任务进度"),("02","核心KPI","关键指标达成分析"),("03","重点工作与问题对策","重点项目推进与问题解决"),("04","下周计划","下周工作安排与规划")]
    for i,(n,tt,de) in enumerate(items):
        y = 1.2 + i * 1.4
        rect(sl, 1.5, y, 0.8, 0.8, C["accent"] if i==0 else C["light_bg"], n, sz=20, bold=True, fc=C["white"] if i==0 else C["primary"])
        tb(sl, 2.6, y, 7, 0.45, tt, 18, True, C["primary"])
        tb(sl, 2.6, y+0.45, 7, 0.3, de, 11, color=C["dark"])
        rect(sl, 2.6, y+0.78, 8, 0.02, C["border"])

# ── Overview (refined table widths) ──
def s_overview(prs, d):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "本周概述", d.get("logo_text","LOGO"))
    summary_bar(sl, ST, d.get("overview_summary",""))
    td = d.get("overview_table",[])
    tt = ST + SHH + 0.1
    if td:
        hd = ["序号","项目名称","项目内容","本周进度","责任人","开始时间","结束时间","进度","状态","备注"]
        rows = [[t.get("seq",""),t.get("project",""),t.get("content",""),t.get("weekly_progress",""),t.get("owner",""),t.get("start",""),t.get("end",""),f"{t.get('progress',0)}%",t.get("status",""),t.get("remark","")] for t in td[:6]]
        th = min(0.32*len(rows)+0.35, 2.8)
        add_table(sl, ML, tt, CW, th, hd, rows, [0.35,1.2,1.8,1.8,0.6,0.7,0.7,0.5,0.6,0.5])
    g = d.get("overview_gantt",{})
    gt = tt + (min(0.32*len(td)+0.35,2.8) if td else 0.3) + 0.15
    gh = H - gt - 0.35
    if g.get("tasks"): gantt(sl, ML, gt, CW, gh, g)

# ── KPI (C: 55/45 left-right split) ──
def s_kpi(prs, d):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "核心KPI看板", d.get("logo_text","LOGO"))
    summary_bar(sl, ST, d.get("kpi_summary",""))
    # 55/45 split: left gets 55% of slide width
    split = W * 0.55
    rect(sl, split, 1.25, 0.02, H-1.6, C["border"])
    lx, lw = ML, split - ML - 0.15
    tb(sl, lx, 1.25, lw, 0.3, d.get("kpi_left_title","指标达成分析"), FONT["section"], True, C["primary"])
    de = d.get("kpi_left_desc","")
    if de:
        rect(sl, lx, 1.6, lw, 0.7, C["light_bg"], "", bc=C["border"], bw=0.5)
        tb(sl, lx+0.1, 1.63, lw-0.2, 0.64, de, FONT["desc"], color=C["dark"])
    ch = d.get("kpi_chart",{})
    if ch.get("series"):
        tb(sl, lx, 2.45, lw, 0.25, f"趋势图: {ch.get('title','')}", FONT["sub"], True, C["primary"])
        add_chart(sl, lx, 2.75, lw, H-3.3, ch)
    rx, rw = split + 0.15, W - MR - (split + 0.15)
    mods = d.get("kpi_right_modules",[])
    if mods:
        mh, mg = 0.82, 0.05
        for i, m in enumerate(mods[:4]):
            y = 1.25 + i*(mh+mg)
            gd = m.get("good",True); bg = C["green"] if gd else C["red"]
            tr = m.get("trend","flat"); ar = "↑" if tr=="up" else ("↓" if tr=="down" else "→")
            rect(sl, rx, y, rw, mh, C["white"], "", bc=C["border"], bw=0.5)
            rect(sl, rx, y, 0.05, mh, bg)
            tb(sl, rx+0.15, y+0.03, rw*0.55, 0.22, m.get("name",""), FONT["kn"], True, C["dark"])
            ac = C["green"] if tr=="up" else (C["red"] if tr=="down" else C["orange"])
            rtb(sl, rx+0.15, y+0.24, rw*0.55, 0.28, [(m.get("value","—"),FONT["kv"],True,bg,"等线"),(f"  {ar}",FONT["kv"]-2,True,ac,"等线")])
            tg = m.get("target","")
            if tg: tb(sl, rx+0.15, y+0.52, rw*0.55, 0.15, f"目标: {tg}", FONT["kt"], color=C["dark"])
    pt = d.get("kpi_right_path_title","达标路径与逻辑")
    py = 1.25 + len(mods)*(mh+mg) + 0.1
    tb(sl, rx, py, rw, 0.25, f"达标路径: {pt}", FONT["sub"], True, C["primary"])
    for i, m in enumerate(mods[:4]):
        p = m.get("path","")
        if p: tb(sl, rx+0.1, py+0.3+i*0.28, rw-0.1, 0.25, f"▸ {m.get('name','')}: {p}", FONT["path"], color=C["dark"])

def s_wp(prs, d):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "重点工作与问题对策", d.get("logo_text","LOGO"))
    summary_bar(sl, ST, d.get("key_works_summary",""))
    kw = d.get("key_works_table",[])
    kt = ST + SHH + 0.08
    if kw:
        hd = ["序号","项目名称","项目内容","本周进度","责任人","开始时间","结束时间","进度","状态","备注"]
        rows = [[t.get("seq",""),t.get("project",""),t.get("content",""),t.get("weekly_progress",""),t.get("owner",""),t.get("start",""),t.get("end",""),f"{t.get('progress',0)}%",t.get("status",""),t.get("remark","")] for t in kw[:4]]
        th = 0.32*len(rows)+0.35
        add_table(sl, ML, kt, CW, th, hd, rows, [0.35,1.2,1.8,1.8,0.6,0.7,0.7,0.5,0.6,0.5])
    py = kt + th + 0.08 if kw else ST + SHH + 0.08
    summary_bar(sl, py, f"问题对策: {d.get('problems_summary','')}")
    pt = d.get("problems_table",[])
    ppt = py + SHH + 0.08
    if pt:
        hd = ["问题描述","影响","解决方案","完成时限","责任人","状态"]
        rows = [[p.get("issue",""),p.get("impact",""),p.get("solution",""),p.get("deadline","—"),p.get("owner",""),p.get("status","")] for p in pt[:3]]
        pth = 0.32*len(rows)+0.35
        add_table(sl, ML, ppt, CW, pth, hd, rows, [2.2,2.0,3.5,1.2,0.8,0.8])

def s_plan(prs, d):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    header(sl, "下周计划", d.get("logo_text","LOGO"))
    summary_bar(sl, ST, d.get("next_plan_summary",""))
    pt = d.get("next_plan_table",[])
    tt = ST + SHH + 0.1
    if pt:
        hd = ["序号","项目名称","项目内容","下周计划","责任人","开始时间","结束时间","进度","状态","备注"]
        rows = [[t.get("seq",""),t.get("project",""),t.get("content",""),t.get("next_plan",""),t.get("owner",""),t.get("start",""),t.get("end",""),f"{t.get('progress',0)}%",t.get("status",""),t.get("remark","")] for t in pt[:6]]
        th = min(0.32*len(rows)+0.35, 2.5)
        add_table(sl, ML, tt, CW, th, hd, rows, [0.35,1.2,1.8,1.8,0.6,0.7,0.7,0.5,0.6,0.5])
    g = d.get("next_plan_gantt",{})
    gt = tt + (min(0.32*len(pt)+0.35,2.5) if pt else 0.3) + 0.15
    gh = H - gt - 0.35
    if g.get("tasks"): gantt(sl, ML, gt, CW, gh, g)

# ══════════════════════════════════════════════════════════════════════════
def generate_pptx(d, out, template_path=None):
    if not PPTX_AVAILABLE:
        return {"error": "pip install python-pptx"}
    if template_path and os.path.exists(template_path):
        # 使用自定义模板
        prs = Presentation(template_path)
        # 清空所有现有幻灯片
        while len(prs.slides) > 0:
            rid = prs.slides._sldIdLst[-1].rId
            prs.part.drop_rel(rid)
            del prs.slides._sldIdLst[-1]
    else:
        # 使用默认空白模板
        prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    s_cover(prs, d)
    s_toc(prs, d)
    s_overview(prs, d)
    s_kpi(prs, d)
    s_wp(prs, d)
    s_plan(prs, d)
    prs.save(out)
    return {"output": out, "slides": len(prs.slides), "template_used": template_path if template_path else "default"}

def main():
    if len(sys.argv) < 3:
        print(json.dumps({"error": "Usage: python chinese_weekly_report.py <data.[json|xlsx]> <output.pptx> [template.pptx]"}))
        sys.exit(1)
    if not os.path.exists(sys.argv[1]):
        print(json.dumps({"error": f"Data file not found: {sys.argv[1]}"}))
        sys.exit(1)
    if Path(sys.argv[2]).suffix.lower() != ".pptx":
        print(json.dumps({"error": "Only .pptx output is supported."}))
        sys.exit(1)
    template_path = sys.argv[3] if len(sys.argv) >=4 and os.path.exists(sys.argv[3]) else None
    try:
        data = load_data(sys.argv[1])
        r = generate_pptx(data, sys.argv[2], template_path)
        print(json.dumps(r, ensure_ascii=False))
        if "error" in r: sys.exit(1)
    except Exception as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        sys.exit(1)

def load_data(p):
    """Load data from JSON or Excel file"""
    suffix = Path(p).suffix.lower()
    if suffix == '.json':
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    elif suffix in ['.xlsx', '.xls']:
        if not PANDAS_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel support, install with: pip install pandas openpyxl")
        d = {}
        # 读取所有工作表
        xls = pd.ExcelFile(p)
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name)
            # 转换为字典格式
            records = df.fillna("").to_dict('records')
            # 映射工作表名到数据字段
            sheet_map = {
                "overview_table": "overview_table",
                "overview_gantt": "overview_gantt",
                "kpi": "kpi_right_modules",
                "kpi_chart": "kpi_chart",
                "key_works": "key_works_table",
                "problems": "problems_table",
                "next_plan": "next_plan_table",
                "next_gantt": "next_plan_gantt",
                "base_info": "base_info"
            }
            if sheet_name in sheet_map:
                d[sheet_map[sheet_name]] = records
            elif sheet_name == "meta":
                # 元数据表，第一列是key，第二列是value
                for _, row in df.iterrows():
                    if len(row) >=2:
                        d[str(row.iloc[0])] = row.iloc[1]
        return d
    else:
        raise ValueError(f"Unsupported file format: {suffix}")

if __name__ == "__main__":
    main()